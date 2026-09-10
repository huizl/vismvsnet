import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "PAPER_DRAFT.md"
OUT = ROOT / "PAPER_DRAFT.docx"


BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
INK = RGBColor(0x0B, 0x25, 0x45)
MUTED = RGBColor(0x55, 0x55, 0x55)
TABLE_HEADER = "F2F4F7"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=80, bottom=80, start=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        elem = tc_mar.find(qn(f"w:{margin_name}"))
        if elem is None:
            elem = OxmlElement(f"w:{margin_name}")
            tc_mar.append(elem)
        elem.set(qn("w:w"), str(value))
        elem.set(qn("w:type"), "dxa")


def set_table_geometry(table, col_widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False

    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_grid = table._tbl.tblGrid
    if tbl_grid is None:
        tbl_grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in col_widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        tbl_grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = col_widths[min(idx, len(col_widths) - 1)]
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_text = OxmlElement("w:t")
    fld_text.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, fld_text, fld_end])


def configure_styles(doc):
    styles = doc.styles

    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    title = styles["Title"]
    title.font.name = "Calibri"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    title.font.size = Pt(20)
    title.font.color.rgb = INK
    title.font.bold = True
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    subtitle.font.size = Pt(11)
    subtitle.font.color.rgb = MUTED
    subtitle.paragraph_format.space_after = Pt(16)

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)

    code_style = styles.add_style("Code Block", 1)
    code_style.font.name = "Consolas"
    code_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    code_style.font.size = Pt(9)
    code_style.paragraph_format.space_before = Pt(4)
    code_style.paragraph_format.space_after = Pt(4)
    code_style.paragraph_format.line_spacing = 1.0


def split_table_row(line):
    parts = [part.strip() for part in line.strip().strip("|").split("|")]
    return parts


def is_table_separator(line):
    stripped = line.strip()
    return bool(re.match(r"^\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?$", stripped))


def parse_markdown(lines):
    blocks = []
    i = 0
    while i < len(lines):
        line = lines[i].rstrip("\n")
        if not line.strip():
            i += 1
            continue
        if line.startswith("```"):
            code = []
            i += 1
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i].rstrip("\n"))
                i += 1
            i += 1
            blocks.append(("code", code))
            continue
        if line.lstrip().startswith("|") and i + 1 < len(lines) and is_table_separator(lines[i + 1]):
            rows = [split_table_row(line)]
            i += 2
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                rows.append(split_table_row(lines[i]))
                i += 1
            blocks.append(("table", rows))
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading:
            blocks.append(("heading", (len(heading.group(1)), heading.group(2).strip())))
            i += 1
            continue
        ordered = re.match(r"^\s*(\d+)\.\s+(.*)$", line)
        if ordered:
            blocks.append(("olist", ordered.group(2).strip()))
            i += 1
            continue
        bullet = re.match(r"^\s*[-*]\s+(.*)$", line)
        if bullet:
            blocks.append(("ulist", bullet.group(1).strip()))
            i += 1
            continue

        para = [line.strip()]
        i += 1
        while i < len(lines):
            nxt = lines[i].rstrip("\n")
            if not nxt.strip():
                break
            if nxt.startswith("```") or re.match(r"^(#{1,6})\s+", nxt):
                break
            if nxt.lstrip().startswith("|") and i + 1 < len(lines) and is_table_separator(lines[i + 1]):
                break
            if re.match(r"^\s*(\d+)\.\s+", nxt) or re.match(r"^\s*[-*]\s+", nxt):
                break
            para.append(nxt.strip())
            i += 1
        blocks.append(("paragraph", " ".join(para)))
    return blocks


def add_formatted_text(paragraph, text):
    token_re = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*)")
    pos = 0
    for match in token_re.finditer(text):
        if match.start() > pos:
            paragraph.add_run(text[pos:match.start()])
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            run.font.size = Pt(9.5)
        elif token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        pos = match.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def clean_title(text):
    return text.replace("`", "")


def build_docx():
    text = SRC.read_text(encoding="utf-8")
    blocks = parse_markdown(text.splitlines())
    doc = Document()

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    configure_styles(doc)

    footer_para = section.footer.paragraphs[0]
    footer_para.style = doc.styles["Normal"]
    footer_para.add_run("PAPER_DRAFT | ")
    add_page_number(footer_para)

    first_h1_used = False
    for block_type, payload in blocks:
        if block_type == "heading":
            level, content = payload
            content = clean_title(content)
            if level == 1 and not first_h1_used:
                first_h1_used = True
                para = doc.add_paragraph(style="Title")
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_formatted_text(para, content)
                sub = doc.add_paragraph("论文初稿 | 基于当前 vismvsnet 代码与实验记录整理", style="Subtitle")
                sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
                doc.add_paragraph()
                continue
            style = "Heading 1" if level <= 2 else "Heading 2" if level == 3 else "Heading 3"
            para = doc.add_paragraph(style=style)
            add_formatted_text(para, content)
        elif block_type == "paragraph":
            para = doc.add_paragraph()
            add_formatted_text(para, payload)
        elif block_type == "olist":
            para = doc.add_paragraph(style="List Number")
            add_formatted_text(para, payload)
        elif block_type == "ulist":
            para = doc.add_paragraph(style="List Bullet")
            add_formatted_text(para, payload)
        elif block_type == "code":
            for line in payload:
                para = doc.add_paragraph(style="Code Block")
                run = para.add_run(line)
                run.font.name = "Consolas"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
                run.font.size = Pt(9)
        elif block_type == "table":
            rows = payload
            if not rows:
                continue
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            widths = [9360 // len(rows[0])] * len(rows[0])
            widths[-1] += 9360 - sum(widths)
            set_table_geometry(table, widths)
            for r_idx, row in enumerate(rows):
                for c_idx, value in enumerate(row):
                    cell = table.cell(r_idx, c_idx)
                    cell.text = ""
                    para = cell.paragraphs[0]
                    para.paragraph_format.space_after = Pt(0)
                    add_formatted_text(para, value)
                    if r_idx == 0:
                        set_cell_shading(cell, TABLE_HEADER)
                        for run in para.runs:
                            run.bold = True
            set_repeat_table_header(table.rows[0])
            doc.add_paragraph()

    doc.core_properties.title = "面向大视差与遮挡场景的几何可见性引导多视图深度估计方法"
    doc.core_properties.subject = "Vis-MVSNet 改进论文初稿"
    doc.core_properties.keywords = "MVS, Vis-MVSNet, occlusion, large disparity, geometric visibility"
    doc.save(OUT)


if __name__ == "__main__":
    build_docx()
